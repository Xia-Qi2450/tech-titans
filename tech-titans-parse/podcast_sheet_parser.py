#!/usr/bin/env python3
"""
podcast_sheet_parser.py

Parses a multi-sheet .xlsx workbook (e.g. Tech Titans Podcast Master Sheet)
and produces two outputs:
  1. A human-readable text report (for people).
  2. A structured JSON file (for other scripts/tools to consume).

Design notes:
  - Parsing is header-driven and generic: it does not hardcode column names,
    so it will keep working if sheets/columns are added, renamed, or
    reordered. Any sheet in the workbook is parsed the same way.
  - Rows that are entirely blank (spreadsheet padding beyond the real data)
    are skipped. Rows that are partially filled (e.g. a placeholder for a
    member slot that hasn't been assigned yet) are kept, and flagged as
    incomplete in the JSON output so downstream tools can filter on that.
  - Member List and Project List have a few columns with controlled
    (enum-style) values. Those are checked against VALIDATION_RULES below;
    unrecognized values aren't dropped or corrected, just flagged with a
    warning (both in the log and as "_validation_warnings" on the record),
    since a typo in the sheet shouldn't make data disappear.
  - Project List's "Project" column encodes two things separated by " - "
    (e.g. "H.E.A.R.T - Honor" -> title "H.E.A.R.T", episode "Honor"). These
    are split out into "project_title" / "episode_title" fields alongside
    the original "project" value.
  - Google Sheets "smart chip" cells (e.g. Deliverable linking to a Doc,
    Student In-Charge linking to a person) export to xlsx as hyperlinked
    cells. Any such cell gets an extra "<column>_link" field with the URL,
    next to that column's own text value.
  - Note-type columns (Script/Recording/Editing Notes) are optional by
    design, so a blank note doesn't count against a record's "_complete"
    flag -- only genuinely-missing required fields do.
  - Logging is used throughout instead of print(), with -v for verbose
    (DEBUG) output and an optional --log-file to persist a run log.

Dependency:
  This script needs `openpyxl` to read .xlsx files (there's no way around
  that for real xlsx parsing without reimplementing the OOXML format).
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from datetime import datetime, date
from pathlib import Path
from typing import Any

try:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
except ImportError:
    print(
        "ERROR: this script requires the 'openpyxl' package.\n"
        "Install it with:  pip install openpyxl",
        file=sys.stderr,
    )
    sys.exit(1)


LOG = logging.getLogger("podcast_sheet_parser")


# --------------------------------------------------------------------------
# Logging setup
# --------------------------------------------------------------------------

def setup_logging(verbose: bool, log_file: Path | None) -> None:
    """Configure the module logger with a console handler and an optional
    file handler."""
    level = logging.DEBUG if verbose else logging.INFO
    fmt = logging.Formatter(
        "%(asctime)s [%(levelname)s] %(message)s", datefmt="%H:%M:%S"
    )

    LOG.setLevel(logging.DEBUG)  # handlers below do their own filtering
    LOG.handlers.clear()

    console = logging.StreamHandler(sys.stdout)
    console.setLevel(level)
    console.setFormatter(fmt)
    LOG.addHandler(console)

    if log_file:
        file_handler = logging.FileHandler(log_file, mode="w", encoding="utf-8")
        file_handler.setLevel(logging.DEBUG)
        file_handler.setFormatter(fmt)
        LOG.addHandler(file_handler)
        LOG.debug("Logging to file: %s", log_file)


# --------------------------------------------------------------------------
# Helpers: normalization
# --------------------------------------------------------------------------

_KEY_SAFE_RE = re.compile(r"[^0-9a-zA-Z]+")


def normalize_key(raw_header: Any, column_index: int, seen: set[str]) -> str:
    """Turn a header cell value into a snake_case JSON-safe key.
    Falls back to 'column_<n>' for blank headers, and de-duplicates
    repeated headers by suffixing _2, _3, ...
    """
    if raw_header is None or str(raw_header).strip() == "":
        key = f"column_{column_index}"
        LOG.warning("Blank header at column %d; using fallback key '%s'", column_index, key)
    else:
        key = _KEY_SAFE_RE.sub("_", str(raw_header).strip()).strip("_").lower()
        if not key:
            key = f"column_{column_index}"

    base_key = key
    suffix = 2
    while key in seen:
        key = f"{base_key}_{suffix}"
        suffix += 1
    if key != base_key:
        LOG.warning("Duplicate header '%s' at column %d; renamed to '%s'", raw_header, column_index, key)

    seen.add(key)
    return key


def normalize_value(value: Any) -> Any:
    """Convert cell values into JSON-friendly, human-friendly forms."""
    if isinstance(value, datetime):
        # Midnight-only datetimes are really just dates in these sheets.
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def row_is_empty(values: list[Any]) -> bool:
    return all(v is None for v in values)


def compute_sheet_keys(sheet_names: list[str]) -> dict[str, str]:
    """Map original sheet names -> snake_case keys, e.g. 'Member List' ->
    'member_list'. Shared between JSON output and validation so both agree
    on the same keys."""
    seen: set[str] = set()
    return {name: normalize_key(name, i, seen) for i, name in enumerate(sheet_names, start=1)}


# --------------------------------------------------------------------------
# Controlled-value validation
# --------------------------------------------------------------------------
# Each sheet maps column keys -> a validation rule:
#   - a list of allowed strings (exact, case-sensitive match)
#   - the literal "DATE" to check the value looks like an ISO date
#   - a compiled regex to check the value matches a pattern
# Missing values (None) are not flagged here -- that's what "_complete"
# is for. This section only checks values that ARE present.

_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}(T\d{2}:\d{2}:\d{2})?$")
_GENERATION_RE = re.compile(r"\d+(st|nd|rd|th)\s+Generation", re.IGNORECASE)

VALIDATION_RULES: dict[str, dict[str, Any]] = {
    "member_list": {
        "status": ["In Service", "Retired", "Yet to Begin"],
        "training_status": ["In Training", "Completed Training", "Haven't Started"],
        "date_of_enlistment": "DATE",
        # Sheet data reads e.g. "Founding Member 1st Generation" rather than
        # a bare "1st Generation" -- validated as a pattern rather than an
        # exact enum for that reason. Flag me if that's not the intent.
        "member_group": _GENERATION_RE,
    },
    "project_list": {
        "priority": ["P0", "P1", "P2", "P3"],
        "script_status": ["Completed", "In Progress", "Not Started"],
        "guest": ["Do Not Have", "Have"],
        "recording_status": ["Completed", "In Progress", "Not Started"],
        "editing_status": ["Completed", "In Progress", "Not Started"],
        "ready_status": ["Ready", "Not Ready"],
        "publish_status": ["Published", "Not Published"],
    },
}


def validate_record(sheet_key: str, record: dict[str, Any]) -> list[str]:
    """Check a record's fields against VALIDATION_RULES for its sheet.
    Returns a list of human-readable warning strings (empty if all clean)."""
    rules = VALIDATION_RULES.get(sheet_key)
    if not rules:
        return []

    warnings: list[str] = []
    for field, rule in rules.items():
        value = record.get(field)
        if value is None:
            continue

        if isinstance(rule, list):
            if value not in rule:
                warnings.append(f"{field}: '{value}' is not a recognized value (expected one of {rule})")
        elif rule == "DATE":
            if not _ISO_DATE_RE.match(str(value)):
                warnings.append(f"{field}: '{value}' does not look like a valid date")
        elif isinstance(rule, re.Pattern):
            if not rule.search(str(value)):
                warnings.append(f"{field}: '{value}' does not match the expected pattern")

    return warnings


# --------------------------------------------------------------------------
# Completeness
# --------------------------------------------------------------------------
# Some columns are known to be optional (blank is a normal, valid state,
# not missing data) -- these are excluded from the "_complete" check.
# Derived fields (things this script adds, like project_title or the
# "<key>_link" hyperlink fields) are excluded too, since they aren't
# source data and would just double up whatever their parent field says.

OPTIONAL_FIELDS: dict[str, set[str]] = {
    "project_list": {"script_notes", "recording_notes", "editing_notes"},
}

DERIVED_FIELDS: dict[str, set[str]] = {
    "project_list": {"project_title", "episode_title"},
}


def compute_completeness(sheet_key: str, record: dict[str, Any]) -> bool:
    skip = OPTIONAL_FIELDS.get(sheet_key, set()) | DERIVED_FIELDS.get(sheet_key, set())
    return all(
        value is not None
        for key, value in record.items()
        if key not in skip and not key.startswith("_") and not key.endswith("_link")
    )


# --------------------------------------------------------------------------
# Project List: derive project_title / episode_title from "Project"
# --------------------------------------------------------------------------

_PROJECT_TITLE_SEP_RE = re.compile(r"\s+-\s+")


def split_project_title(record: dict[str, Any]) -> dict[str, Any]:
    """Split the 'project' field on ' - ' into project_title / episode_title,
    e.g. 'H.E.A.R.T - Honor' -> title='H.E.A.R.T', episode='Honor'.
    Returns a new dict with the derived fields inserted right after
    'project', preserving the rest of the field order."""
    project_value = record.get("project")
    title: str | None = None
    episode: str | None = None

    if project_value:
        parts = _PROJECT_TITLE_SEP_RE.split(project_value, maxsplit=1)
        if len(parts) == 2:
            title, episode = parts[0].strip(), parts[1].strip()
        else:
            title = project_value.strip()
            LOG.debug("Project '%s' has no ' - ' separator; treating whole string as the title", project_value)

    new_record: dict[str, Any] = {}
    for key, value in record.items():
        new_record[key] = value
        if key == "project":
            new_record["project_title"] = title
            new_record["episode_title"] = episode
    return new_record


def postprocess_records(data: dict[str, list[dict[str, Any]]], sheet_key_map: dict[str, str]) -> None:
    """Apply sheet-specific derived fields, completeness, and controlled-value
    validation in place, after parsing."""
    for sheet_name, records in data.items():
        sheet_key = sheet_key_map[sheet_name]

        for i, record in enumerate(records):
            record["_complete"] = compute_completeness(sheet_key, record)

            if sheet_key == "project_list":
                record = split_project_title(record)
                records[i] = record

            warnings = validate_record(sheet_key, record)
            record["_validation_warnings"] = warnings
            if warnings:
                identifier = record.get("name") or record.get("project") or f"record #{i + 1}"
                for warning in warnings:
                    LOG.warning("Sheet '%s' [%s]: %s", sheet_name, identifier, warning)

        incomplete = sum(1 for r in records if not r["_complete"])
        if incomplete:
            LOG.info("Sheet '%s': %d record(s) are missing required (non-optional) fields", sheet_name, incomplete)


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------

def parse_sheet(ws) -> list[dict[str, Any]]:
    """Parse a single worksheet into a list of record dicts, using row 1
    as the header row.

    Any cell that carries a hyperlink (Google Sheets "smart chips" -- linked
    docs, people, etc. -- come through the xlsx export as a hyperlinked
    cell) gets an extra "<key>_link" field with the link target, right next
    to the cell's own value. This isn't specific to any one column: whatever
    column happens to have a hyperlink gets one.
    """
    rows_iter = ws.iter_rows()
    try:
        header_row = next(rows_iter)
    except StopIteration:
        LOG.warning("Sheet '%s' is completely empty (no header row)", ws.title)
        return []

    seen_keys: set[str] = set()
    keys = [
        normalize_key(cell.value, idx, seen_keys) for idx, cell in enumerate(header_row, start=1)
    ]

    records: list[dict[str, Any]] = []
    skipped_blank = 0
    links_found = 0

    for row in rows_iter:
        if row_is_empty([cell.value for cell in row]):
            skipped_blank += 1
            continue

        # Pad/truncate row to header length defensively (ragged rows happen).
        cells = list(row) + [None] * (len(keys) - len(row))
        cells = cells[: len(keys)]

        record: dict[str, Any] = {}
        for key, cell in zip(keys, cells):
            record[key] = normalize_value(cell.value) if cell is not None else None
            link = cell.hyperlink.target if (cell is not None and cell.hyperlink) else None
            if link:
                record[f"{key}_link"] = link
                links_found += 1

        records.append(record)

    LOG.info(
        "Sheet '%s': parsed %d record(s), skipped %d blank row(s)",
        ws.title, len(records), skipped_blank,
    )
    if links_found:
        LOG.info("Sheet '%s': found %d hyperlinked cell(s)", ws.title, links_found)

    return records


def parse_workbook(path: Path) -> dict[str, list[dict[str, Any]]]:
    LOG.info("Loading workbook: %s", path)
    try:
        # read_only=False (the default) is required here: openpyxl's
        # read-only mode returns cells with no .hyperlink attribute at all,
        # which would silently lose the smart-chip links. These sheets are
        # small enough that the extra memory doesn't matter.
        wb = openpyxl.load_workbook(path, data_only=True)
    except InvalidFileException as exc:
        LOG.error("Not a valid .xlsx file: %s", exc)
        raise
    except FileNotFoundError:
        LOG.error("File not found: %s", path)
        raise

    LOG.debug("Sheets found: %s", wb.sheetnames)

    data: dict[str, list[dict[str, Any]]] = {}
    for sheet_name in wb.sheetnames:
        LOG.debug("Parsing sheet '%s'", sheet_name)
        data[sheet_name] = parse_sheet(wb[sheet_name])

    wb.close()
    return data


# --------------------------------------------------------------------------
# Output: human-readable
# --------------------------------------------------------------------------

def build_human_readable(data: dict[str, list[dict[str, Any]]], source: Path) -> str:
    lines: list[str] = []
    lines.append("=" * 70)
    lines.append(f"REPORT: {source.name}")
    lines.append(f"Generated: {datetime.now().isoformat(timespec='seconds')}")
    lines.append("=" * 70)
    lines.append("")
    lines.append("Summary:")
    for sheet_name, records in data.items():
        incomplete = sum(1 for r in records if not r["_complete"])
        lines.append(f"  - {sheet_name}: {len(records)} record(s), {incomplete} incomplete")
    lines.append("")

    for sheet_name, records in data.items():
        lines.append("-" * 70)
        lines.append(f"SHEET: {sheet_name}")
        lines.append("-" * 70)

        if not records:
            lines.append("  (no data)")
            lines.append("")
            continue

        for i, record in enumerate(records, start=1):
            flag = "" if record["_complete"] else "  [INCOMPLETE]"
            lines.append(f"[{i}]{flag}")
            for key, value in record.items():
                if key.startswith("_"):
                    continue  # metadata fields (_complete, _validation_warnings) shown separately
                display_key = key.replace("_", " ").title()
                display_value = "(empty)" if value is None else value
                lines.append(f"    {display_key}: {display_value}")

            warnings = record.get("_validation_warnings") or []
            if warnings:
                lines.append("    \u26a0 Validation warnings:")
                for warning in warnings:
                    lines.append(f"        - {warning}")

            lines.append("")

    return "\n".join(lines)


# --------------------------------------------------------------------------
# Output: JSON
# --------------------------------------------------------------------------

def build_json_payload(
    data: dict[str, list[dict[str, Any]]], source: Path, sheet_key_map: dict[str, str]
) -> dict[str, Any]:
    total_warnings = sum(
        1 for records in data.values() for r in records if r.get("_validation_warnings")
    )

    return {
        "metadata": {
            "source_file": source.name,
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "sheet_record_counts": {name: len(records) for name, records in data.items()},
            "records_with_validation_warnings": total_warnings,
        },
        "data": {
            sheet_key_map[name]: {
                "sheet_name": name,
                "records": records,
            }
            for name, records in data.items()
        },
    }


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Parse an .xlsx workbook into a human-readable report and JSON."
    )
    parser.add_argument("input", type=Path, help="Path to the source .xlsx file")
    parser.add_argument(
        "-o", "--output-dir", type=Path, default=Path("./output"),
        help="Directory to write outputs to (default: ./output)",
    )
    parser.add_argument("--json-name", type=str, default=None, help="Filename for the JSON output")
    parser.add_argument("--text-name", type=str, default=None, help="Filename for the text report")
    parser.add_argument("--json-indent", type=int, default=2, help="JSON indent width, 0 for compact (default: 2)")
    parser.add_argument("--stdout", action="store_true", help="Also print the human-readable report to stdout")
    parser.add_argument("-v", "--verbose", action="store_true", help="Enable debug logging")
    parser.add_argument("--log-file", type=Path, default=None, help="Also write the run log to this file")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    setup_logging(args.verbose, args.log_file)

    if not args.input.exists():
        LOG.error("Input file does not exist: %s", args.input)
        return 1

    try:
        data = parse_workbook(args.input)
    except Exception:
        LOG.exception("Failed to parse workbook")
        return 1

    sheet_key_map = compute_sheet_keys(list(data.keys()))
    postprocess_records(data, sheet_key_map)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / (args.json_name or f"{args.input.stem}.json")
    text_path = args.output_dir / (args.text_name or f"{args.input.stem}.txt")

    report_text = build_human_readable(data, args.input)
    text_path.write_text(report_text, encoding="utf-8")
    LOG.info("Wrote human-readable report: %s", text_path)

    json_payload = build_json_payload(data, args.input, sheet_key_map)
    indent = args.json_indent if args.json_indent > 0 else None
    json_path.write_text(json.dumps(json_payload, indent=indent, ensure_ascii=False), encoding="utf-8")
    LOG.info("Wrote JSON output: %s", json_path)

    if args.stdout:
        print()
        print(report_text)

    LOG.info("Done.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
